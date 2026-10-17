#!/usr/bin/env python3
"""Emergency squawks worldwide, from the AeroSignals bot on Bluesky, replayed through flight_watch.

AeroSignals (https://bsky.app/profile/spotterintel.bsky.social, since 11 Apr 2026) posts every
airliner squawking 7700 / 7600 / 7500 a few minutes after it starts, as a fixed template:

    Ryanair flight squawks 7700.
    Ryanair FR6789 | 737-800
    Ronchi De Legionari (TRS) → Beauvais (BVA)

    Squawked 7700. Monitoring.

and later replies with the outcome ("Update: ... diverted to Venice (VCE), has landed safely ...").
Its posts are read through Bluesky's public API (no login). For each squawk post:

1. the callsign: the airline's ICAO code (the name in the post picks among airlines sharing an IATA
   code) + the flight number, or a callsign whose route in the VRS standing data is the post's route
   (Ryanair flies FR1386 as an alphanumeric callsign);
2. the aircraft: the callsigns sent in that half hour, from ADS-B Exchange's heatmap (17 MB a file;
   only a small callsign index is kept), then the day's trace of each candidate: the one squawking
   the code before the post is it;
3. the flight: the trace's leg around the squawk (airborne between two ground stops), replayed
   through the real Monitor by tools/track_alerts.py (the procedure the map page runs with `analyze`),
   watched from the origin airport and from the destination as if each were ours (the flight followed
   by its route beyond the 250 nm circle), and from the squawk's point (every report local).

Results go to squawks/ (git-ignored): posts.jsonl (the bot's posts), cases.jsonl (one line a post:
what was found and our alerts), heatmaps/ (callsign indexes), traces/ and squawk_cases.xlsx. A run
reads only new posts and does only new cases; posts younger than --min-age-h wait for a later run
(the flight may still be in the air). Nothing is asked of adsb.lol (the monitor's feed): the
heatmaps and traces are ADS-B Exchange's, one request at a time with a pause between them.

  .venv/bin/python tools/squawk_cases.py              # update everything, write the spreadsheet
  .venv/bin/python tools/squawk_cases.py --limit 20   # at most 20 new cases this run
  .venv/bin/python tools/squawk_cases.py --redo       # replay every found case again (code changed)

Needs openpyxl (`.venv/bin/pip install openpyxl`) and the local standing-data/ checkout.
"""
from __future__ import annotations

import argparse
import csv
import glob
import json
import logging
import os
import re
import sys
import tempfile
import time
import urllib.error
import urllib.parse
import urllib.request
from datetime import datetime, timedelta, timezone
from zoneinfo import ZoneInfo

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, ".."))
sys.path.insert(0, os.path.join(HERE, "..", "tests"))
import flight_watch as fw  # noqa: E402
import capture_incident as ci  # noqa: E402
import track_alerts  # noqa: E402

ACTOR = "spotterintel.bsky.social"
FEED = "https://public.api.bsky.app/xrpc/app.bsky.feed.getAuthorFeed"
OUT = os.path.join(HERE, "..", "squawks")
PAUSE = 1.5  # s between requests to ADS-B Exchange / Bluesky
SQUAWK_RE = re.compile(r"Squawked (7[567]00)")
UTC = timezone.utc
ISRAEL = ZoneInfo("Asia/Jerusalem")
log = logging.getLogger("squawks")


# ---------------------------------------------------------------- the bot's posts

def fetch_posts(path: str) -> list[dict]:
    """All posts and replies of the bot, newest first; only pages newer than the file are read."""
    old = []
    if os.path.exists(path):
        with open(path, encoding="utf-8") as f:
            old = [json.loads(line) for line in f if line.strip()]
    known = {p["uri"] for p in old}
    new, cursor = [], None
    while True:
        q = {"actor": ACTOR, "limit": 100, "filter": "posts_with_replies", **({"cursor": cursor} if cursor else {})}
        with urllib.request.urlopen(f"{FEED}?{urllib.parse.urlencode(q)}", timeout=60) as r:
            page = json.load(r)
        fresh = 0
        for item in page.get("feed", []):
            p = item["post"]
            if p["author"]["handle"] != ACTOR or p["uri"] in known:
                continue
            rec = p["record"]
            reply = rec.get("reply") or {}
            new.append({"uri": p["uri"], "t": rec["createdAt"], "text": rec.get("text", ""),
                        "root": (reply.get("root") or {}).get("uri")})
            known.add(p["uri"])
            fresh += 1
        cursor = page.get("cursor")
        if not cursor or not page.get("feed") or (old and not fresh):
            break
        time.sleep(PAUSE)
    posts = sorted(new + old, key=lambda p: p["t"], reverse=True)
    if new:
        with open(path, "w", encoding="utf-8") as f:
            for p in posts:
                f.write(json.dumps(p, ensure_ascii=False) + "\n")
    log.info("posts: %d (%d new)", len(posts), len(new))
    return posts


def parse_case(p: dict) -> dict | None:
    """The squawk post's fields, or None for other posts (deliveries, diversions, news)."""
    m = SQUAWK_RE.search(p["text"])
    if not m or p.get("root"):
        return None
    lines = [x.strip() for x in p["text"].splitlines() if x.strip()]
    c = {"uri": p["uri"], "post_t": p["t"], "squawk": m.group(1), "headline": lines[0]}
    for line in lines[1:]:
        if "|" in line and "airline" not in c:
            who, c["type_name"] = (x.strip() for x in line.split("|", 1))
            fm = re.search(r"\s([A-Z0-9]{2})\s?(\d{1,4}[A-Z]?)$", who)
            if fm:
                c["airline"], c["iata"], c["number"] = who[:fm.start()].strip(), fm.group(1), fm.group(2)
            else:
                c["airline"], c["flight_text"] = who, who
        elif "→" in line:
            a, b = line.split("→", 1)
            c["from"], c["to"] = _code(a), _code(b)
            c["from_name"], c["to_name"] = a.strip(), b.strip()
        elif line.startswith("From "):
            c["from"], c["from_name"] = _code(line), line[5:].strip()
        elif line.startswith("To "):
            c["to"], c["to_name"] = _code(line), line[3:].strip()
    return c


def _code(s: str) -> str | None:
    m = re.search(r"\(([A-Z0-9]{3})\)", s)
    return m.group(1) if m else None


def outcomes(posts: list[dict]) -> dict:
    """Root post uri -> the bot's last word on it ("diverted to Venice (VCE), has landed safely ...")."""
    out = {}
    for p in sorted(posts, key=lambda p: p["t"]):
        if p.get("root") and p["text"].startswith("Update:"):
            out[p["root"]] = p["text"][len("Update:"):].strip()
    return out


# ---------------------------------------------------------------- reference data

class Reference:
    """Airports by IATA code and airlines by IATA code, from the local standing-data checkout."""

    def __init__(self, source: str):
        self.airports = {}  # IATA -> (ICAO, lat, lon, name)
        for name in glob.glob(os.path.join(source, "airports", "schema-01", "*", "*.csv")):
            with open(name, encoding="utf-8-sig") as f:
                for row in csv.reader(f):
                    try:
                        if row[3] and row[2]:
                            self.airports.setdefault(row[3], (row[2], float(row[6]), float(row[7]), row[1]))
                    except (ValueError, IndexError):
                        continue
        self.airlines = {}  # IATA -> [(ICAO, name)]
        with open(os.path.join(source, "airlines", "schema-01", "airlines.csv"), encoding="utf-8-sig") as f:
            for row in csv.reader(f):
                if len(row) > 3 and row[3] and row[2] and row[0] != "Code":
                    self.airlines.setdefault(row[3], []).append((row[2], row[1]))
        self.standing = fw.StandingData(None, source)
        if not self.airports:
            raise SystemExit(f"no airports under {source}: clone the standing data there "
                             "(git clone https://github.com/vradarserver/standing-data)")

    def airline_icao(self, iata: str, name: str) -> list[str]:
        """ICAO codes for an IATA airline code, the name's best match first."""
        cands = self.airlines.get(iata, [])
        words = set(re.findall(r"[a-z]+", (name or "").lower())) - {"airlines", "airline", "air", "airways"}
        cands = sorted(cands, key=lambda c: -len(words & set(re.findall(r"[a-z]+", c[1].lower()))))
        out = [c[0] for c in cands]
        if iata in fw.AIRLINE_ICAO:
            out = [fw.AIRLINE_ICAO[iata]] + [x for x in out if x != fw.AIRLINE_ICAO[iata]]
        return out


# ---------------------------------------------------------------- ADS-B Exchange history

def get(url: str) -> bytes:
    time.sleep(PAUSE)
    return ci.get(url)


def callsign_index(day: datetime, n: int, cache: str) -> dict:
    """Callsign -> [hex] in ADS-B Exchange's half-hour heatmap n of `day` (index cached, file not)."""
    path = os.path.join(cache, f"{day:%Y%m%d}_{n:02d}.json")
    if os.path.exists(path):
        with open(path) as f:
            return json.load(f)
    seen: dict = {}
    with tempfile.TemporaryDirectory() as tmp:
        time.sleep(PAUSE)
        for _ in ci.heatmap(day, n, tmp, seen):
            pass
    index = {cs: sorted(h) for cs, h in seen.items()}
    os.makedirs(cache, exist_ok=True)
    with open(path, "w") as f:
        json.dump(index, f, separators=(",", ":"))
    return index


def day_trace(hexid: str, day: datetime, cache: str) -> dict | None:
    path = os.path.join(cache, f"{day:%Y%m%d}_{hexid}.json")
    if not os.path.exists(path):
        try:
            data = get(f"{ci.BASE}/{day:%Y/%m/%d}/traces/{hexid[-2:]}/trace_full_{hexid}.json")
        except urllib.error.HTTPError as e:
            if e.code != 404:
                raise
            data = b"{}"
        os.makedirs(cache, exist_ok=True)
        with open(path, "wb") as f:
            f.write(data)
    with open(path, "rb") as f:
        tr = json.loads(f.read() or b"{}")
    return tr if tr.get("trace") else None


def joined_trace(hexid: str, days: list[datetime], cache: str) -> dict | None:
    """Several UTC days of one aircraft as one trace (rows rebased on the first day's timestamp)."""
    out = None
    for day in days:
        tr = day_trace(hexid, day, cache)
        if not tr:
            continue
        if out is None:
            out = {**tr, "trace": list(tr["trace"])}
            continue
        shift = tr["timestamp"] - out["timestamp"]
        last = out["trace"][-1][0] if out["trace"] else -1e18
        out["trace"] += [[p[0] + shift, *p[1:]] for p in tr["trace"] if p[0] + shift > last]
    return out


def squawk_start(tr: dict, code: str, lo: float, hi: float) -> tuple[float, list] | None:
    """(time, row) of the first report squawking `code` in [lo, hi], following an earlier squawk
    of it back to where it started."""
    first, prev = None, None
    for p in tr["trace"]:
        t = tr["timestamp"] + p[0]
        ex = p[8] if len(p) > 8 and isinstance(p[8], dict) else None
        if ex is None:
            continue
        sq = ex.get("squawk")
        if sq == code:
            if prev != code:
                first = (t, p)
            if lo <= t <= hi:
                return first
        prev = sq
        if t > hi:
            break
    return None


def candidates(c: dict, ref: Reference, index: dict) -> list[str]:
    """Callsigns of the post's flight among those in the heatmap index, most likely first."""
    if "iata" not in c:
        return []
    out = []
    num = c["number"].lstrip("0")
    prefixes = ref.airline_icao(c["iata"], c.get("airline", ""))
    want = [ref.airports[x][0] for x in (c.get("from"), c.get("to")) if x in ref.airports]
    for icao in prefixes:
        out += [cs for cs in index if cs.startswith(icao) and cs[len(icao):].lstrip("0") == num]
    for icao in prefixes:  # alphanumeric callsigns: the route data's route is the post's
        for cs in index:
            if cs.startswith(icao) and cs not in out and len(want) == 2:
                try:
                    r = ref.standing.route(cs)
                except Exception:  # noqa: BLE001 - a broken file is no candidate
                    r = None
                if r and r[0] == want[0] and r[-1] == want[1]:
                    out.append(cs)
    return out


def find_flight(c: dict, ref: Reference, cache: str) -> dict:
    """Callsign, hex and trace of the post's flight, or {"status": why not}."""
    post = datetime.fromisoformat(c["post_t"].replace("Z", "+00:00"))
    lo, hi = (post - timedelta(minutes=90)).timestamp(), (post + timedelta(minutes=5)).timestamp()
    tried = []
    for back in (5, 40):  # the half hour before the post, then the one before that
        t = post - timedelta(minutes=back)
        day = t.replace(hour=0, minute=0, second=0, microsecond=0)
        n = (t.hour * 60 + t.minute) // 30
        try:
            index = callsign_index(day, n, os.path.join(cache, "heatmaps"))
        except urllib.error.HTTPError as e:
            return {"status": f"heatmap {day:%Y-%m-%d} {n}: HTTP {e.code}"}
        for cs in candidates(c, ref, index):
            for hexid in index[cs]:
                if (cs, hexid) in tried:
                    continue
                tried.append((cs, hexid))
                days = [day - timedelta(days=1), day, day + timedelta(days=1)]
                days = [d for d in days if d.timestamp() < time.time()]
                tr = joined_trace(hexid, [day], os.path.join(cache, "traces"))
                if not tr:
                    continue
                hit = squawk_start(tr, c["squawk"], lo, hi)
                if hit:
                    tr = joined_trace(hexid, days, os.path.join(cache, "traces"))
                    return {"status": "found", "callsign": cs, "hex": hexid, "trace": tr,
                            "squawk_t": hit[0], "squawk_row": hit[1]}
    if not tried:
        return {"status": "callsign not found" if "iata" in c else "no flight number in the post"}
    return {"status": f"no {c['squawk']} in the traces of " + ", ".join(f"{cs} {h}" for cs, h in tried)}


# ---------------------------------------------------------------- the replay

def phase(row: list) -> str:
    alt, rate = row[3], row[7]
    if alt == "ground":
        return "ground"
    if rate is not None and rate >= 500:
        return "climb"
    if rate is not None and rate <= -500:
        return "descent"
    return "level"


def do_case(c: dict, ref: Reference, cache: str) -> dict:
    found = find_flight(c, ref, cache)
    rec = {k: v for k, v in c.items()}
    rec["status"] = found["status"]
    if found["status"] != "found":
        return rec
    tr, st = found["trace"], found["squawk_t"]
    row = found["squawk_row"]
    rec.update(callsign=found["callsign"], hex=found["hex"], reg=tr.get("r"), actype=tr.get("t"),
               desc=tr.get("desc"), squawk_t=st, squawk_alt=row[3], squawk_phase=phase(row),
               squawk_pos=[round(row[1], 4), round(row[2], 4)])
    ends = [{"code": x, "icao": ref.airports[x][0], "lat": ref.airports[x][1], "lon": ref.airports[x][2],
             "name": ref.airports[x][3]} for x in (c.get("from"), c.get("to")) if x in ref.airports]
    got = track_alerts.analyze([tr], found["hex"], st, ends, c.get("iata", "") + c.get("number", ""))
    if "error" in got:
        rec["status"] = got["error"]
        return rec
    rec["leg"], rec["alerts"] = got["leg"], got["alerts"]
    end = got["last"]
    if end[2] in (0, "ground"):
        near = ref.standing.airfield_near(end[0], end[1])
        rec["landed"] = near[0] if near and near[1] < 5 else f"{end[0]:.2f},{end[1]:.2f}"
    else:
        rec["landed"] = None
        rec["last_heard"] = [round(end[0], 3), round(end[1], 3), end[2]]
    return rec


# ---------------------------------------------------------------- spreadsheet

def fmt_t(t: float | None) -> str:
    return time.strftime("%H:%M:%S", time.gmtime(t)) if t else ""


def write_xlsx(cases: list[dict], outs: dict, path: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    skip = set(fw.parse_args([]).telegram_skip.split(","))
    wb = Workbook()
    ws = wb.active
    ws.title = "squawks"
    cols = ["date (UTC)", "post (UTC)", "squawk", "airline", "flight", "callsign", "hex", "reg", "type",
            "from", "to", "outcome (bot)", "landed (trace)", "squawk start (UTC)", "bot lag min",
            "squawk alt ft", "phase", "our alerts", "other signs", "first other sign (min vs squawk)",
            "on Telegram (other)", "bluesky post", "our map", "map + analysis", "ADSBx replay", "status"]
    ws.append(cols)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    link_font = Font(color="0563C1", underline="single")
    for c in sorted(cases, key=lambda c: c["post_t"], reverse=True):
        st = c.get("squawk_t")
        alerts = c.get("alerts") or []
        other = [a for a in alerts if a["kind"] != "EMERGENCY"]
        rel = lambda a: (a["time"] - st) / 60  # noqa: E731
        first_other = round(rel(other[0]), 1) if other and st else ""
        post = datetime.fromisoformat(c["post_t"].replace("Z", "+00:00"))
        alt = c.get("squawk_alt")
        tg = sorted({a["kind"] for a in other if a["kind"] not in skip})
        # the map of our first alert, as it would have been posted, with all the flight's alerts (a=)
        mapq = None
        if c.get("hex") and st:
            t0 = int(alerts[0]["time"]) if alerts else int(st)
            mapq = {"hex": c["hex"], "t": t0}
            if alerts:
                mapq["a"] = fw.map_alerts([(a["kind"], a["time"]) for a in alerts], t0)
        ws.append([
            post.strftime("%Y-%m-%d"), post.strftime("%H:%M"), c["squawk"], c.get("airline", ""),
            (c.get("iata", "") + c.get("number", "")) or c.get("flight_text", ""), c.get("callsign", ""),
            c.get("hex", ""), c.get("reg", ""), c.get("actype", ""), c.get("from", ""), c.get("to", ""),
            outs.get(c["uri"], ""), c.get("landed") or ("last heard " + ",".join(map(str, c["last_heard"]))
                                                        if c.get("last_heard") else ""),
            fmt_t(st), round((post.timestamp() - st) / 60, 1) if st else "",
            alt if alt != "ground" else 0, c.get("squawk_phase", ""),
            "; ".join(f"{a['kind']} {fmt_t(a['time'])[:5]} ({rel(a):+.0f})" for a in alerts) if st else "",
            ", ".join(dict.fromkeys(a["kind"] for a in other)), first_other, ", ".join(tg),
            "post", "map" if mapq else "", "analyse" if mapq else "", "replay" if mapq else "", c["status"]])
        r = ws.max_row
        rkey = c["uri"].rsplit("/", 1)[1]
        links = {22: f"https://bsky.app/profile/{ACTOR}/post/{rkey}"}
        if mapq:
            links[23] = fw.VIEWER_URL + "?" + urllib.parse.urlencode(mapq, safe="~")
            links[24] = links[23] + "&analyze"
            links[25] = (f"https://globe.adsbexchange.com/?icao={c['hex']}&showTrace="
                         f"{time.strftime('%Y-%m-%d', time.gmtime(mapq['t']))}&timestamp={mapq['t']}")
        for col, url in links.items():
            cell = ws.cell(row=r, column=col)
            cell.hyperlink, cell.font = url, link_font
    widths = [11, 7, 7, 22, 9, 10, 8, 9, 6, 5, 5, 50, 12, 10, 6, 8, 8, 60, 30, 10, 20, 6, 6, 8, 7, 30]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    s = wb.create_sheet("summary")
    found = [c for c in cases if c.get("status", "").startswith("found")]
    s.append(["posts with a squawk", len(cases)])
    s.append(["flight found and replayed", sum(1 for c in found if c.get("alerts") is not None)])
    s.append(["not found", len(cases) - len(found)])
    s.append([])
    s.append(["alert kind (other than EMERGENCY)", "flights", "first before the squawk"])
    kinds: dict = {}
    for c in found:
        for k, a in {a["kind"]: a for a in reversed(c.get("alerts") or [])}.items():
            if k == "EMERGENCY":
                continue
            before = a["time"] < c["squawk_t"]
            n = kinds.setdefault(k, [0, 0])
            n[0] += 1
            n[1] += before
    for k, (n, b) in sorted(kinds.items(), key=lambda x: -x[1][0]):
        s.append([k, n, b])
    s.append([])
    s.append(["flights with no alert but EMERGENCY",
              sum(1 for c in found if c.get("alerts") is not None and all(a["kind"] == "EMERGENCY" for a in c["alerts"]))])
    s.column_dimensions["A"].width = 40
    wb.save(path)


# ---------------------------------------------------------------- main

def main() -> None:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--out", default=OUT, help="results folder (git-ignored)")
    p.add_argument("--standing-data", default=os.path.join(HERE, "..", "standing-data"))
    p.add_argument("--limit", type=int, default=0, help="at most this many new cases this run (0: all)")
    p.add_argument("--min-age-h", type=float, default=6, help="posts younger than this wait for a later run")
    p.add_argument("--redo", action="store_true", help="replay the found cases again (their traces are cached)")
    p.add_argument("--retry", action="store_true", help="look again for cases not found before")
    p.add_argument("--match", help="only posts whose text contains this (e.g. FZ1073), done again")
    p.add_argument("--no-fetch", action="store_true", help="use the posts already downloaded")
    a = p.parse_args()
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(message)s", datefmt="%H:%M:%S")
    logging.getLogger("flight_watch").setLevel(logging.ERROR)
    os.makedirs(a.out, exist_ok=True)
    posts_path = os.path.join(a.out, "posts.jsonl")
    posts = fetch_posts(posts_path) if not a.no_fetch else [json.loads(x) for x in open(posts_path, encoding="utf-8")]
    ref = Reference(a.standing_data)
    cases_path = os.path.join(a.out, "cases.jsonl")
    done = {}
    if os.path.exists(cases_path):
        with open(cases_path, encoding="utf-8") as f:
            for line in f:
                r = json.loads(line)
                done[r["uri"]] = r
    todo = []
    cutoff = time.time() - a.min_age_h * 3600
    for post in sorted(posts, key=lambda p: p["t"]):
        c = parse_case(post)
        if not c or datetime.fromisoformat(c["post_t"].replace("Z", "+00:00")).timestamp() > cutoff:
            continue
        old = done.get(c["uri"])
        if a.match:
            if a.match in post["text"]:
                todo.append(c)
            continue
        if old is None or (a.retry and not old["status"].startswith("found")) or (a.redo and old.get("hex")):
            todo.append(c)
    if a.limit:
        todo = todo[:a.limit]
    log.info("cases: %d done, %d to do", len(done), len(todo))
    for i, c in enumerate(todo, 1):
        try:
            rec = do_case(c, ref, a.out)
        except Exception as e:  # noqa: BLE001 - one bad case must not stop the rest
            log.exception("case %s failed", c["uri"])
            rec = {**c, "status": f"error: {type(e).__name__}: {e}"}
        done[c["uri"]] = rec
        kinds = ", ".join(dict.fromkeys(x["kind"] for x in rec.get("alerts") or [])) or "-"
        log.info("%d/%d %s %s %s %s: %s | %s", i, len(todo), c["post_t"][:16], c["squawk"],
                 c.get("iata", "") + c.get("number", ""), rec.get("callsign", ""), rec["status"], kinds)
        with open(cases_path, "w", encoding="utf-8") as f:  # rewritten after each case: safe to stop any time
            for r in done.values():
                f.write(json.dumps(r, ensure_ascii=False) + "\n")
    xlsx = os.path.join(a.out, "squawk_cases.xlsx")
    write_xlsx(list(done.values()), outcomes(posts), xlsx)
    log.info("wrote %s (%d cases)", xlsx, len(done))


if __name__ == "__main__":
    main()
